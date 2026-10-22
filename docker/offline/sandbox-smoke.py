"""在临时目录中验证离线办公依赖、中文转换和非 root 写入边界。"""

from __future__ import annotations

import json
import os
import subprocess
import sys
import tempfile
import zipfile
from pathlib import Path

from docx import Document
from docxtpl import DocxTemplate
from openpyxl import Workbook, load_workbook
from PIL import Image, ImageDraw, ImageFont
from pptx import Presentation
from pypdf import PdfReader
from reportlab.pdfgen import canvas


def _run(command: list[str], *, cwd: Path, timeout: int = 90) -> str:
    """运行检查命令并在失败时保留可诊断错误。"""
    result = subprocess.run(
        command,
        cwd=cwd,
        check=False,
        capture_output=True,
        text=True,
        timeout=timeout,
        env={**os.environ, "HOME": str(cwd), "DOTNET_CLI_HOME": str(cwd / ".dotnet")},
    )
    if result.returncode:
        raise RuntimeError(
            f"命令失败 ({result.returncode}): {command!r}; "
            f"stdout={result.stdout[-2000:]!r}; stderr={result.stderr[-4000:]!r}"
        )
    return result.stdout.strip()


def _write_dotnet_docx(root: Path) -> bool:
    """离线还原预缓存的 OpenXML 包并由 .NET 实际生成 DOCX。"""
    package_cache = Path(os.environ.get("NUGET_PACKAGES", "/opt/yuxi-nuget"))
    package_dir = package_cache / "documentformat.openxml" / "3.3.0"
    if not package_dir.is_dir():
        return False

    project = root / "dotnet-openxml"
    project.mkdir()
    (project / "OpenXmlSmoke.csproj").write_text(
        """<Project Sdk=\"Microsoft.NET.Sdk\">
  <PropertyGroup><OutputType>Exe</OutputType><TargetFramework>net8.0</TargetFramework></PropertyGroup>
  <ItemGroup><PackageReference Include=\"DocumentFormat.OpenXml\" Version=\"3.3.0\" /></ItemGroup>
</Project>
""",
        encoding="utf-8",
    )
    (project / "Program.cs").write_text(
        """using DocumentFormat.OpenXml;
using DocumentFormat.OpenXml.Packaging;
using DocumentFormat.OpenXml.Wordprocessing;
using var document = WordprocessingDocument.Create(args[0], WordprocessingDocumentType.Document);
var part = document.AddMainDocumentPart();
part.Document = new Document(new Body(new Paragraph(new Run(new Text(\"离线 OpenXML 验证\")))));
part.Document.Save();
""",
        encoding="utf-8",
    )
    output = root / "dotnet-openxml.docx"
    _run(
        [
            "dotnet",
            "restore",
            str(project / "OpenXmlSmoke.csproj"),
            "--source",
            str(package_cache),
            "--ignore-failed-sources",
        ],
        cwd=root,
        timeout=180,
    )
    _run(
        [
            "dotnet",
            "run",
            "--no-restore",
            "--project",
            str(project / "OpenXmlSmoke.csproj"),
            "--",
            str(output),
        ],
        cwd=root,
        timeout=180,
    )
    return output.is_file() and zipfile.is_zipfile(output)


def _write_node_office(root: Path) -> dict[str, bool]:
    """验证 Node 包可直接导入，并能生成有效 Office ZIP 产物。"""
    script = r"""
const fs = require('fs');
const { Document, Packer, Paragraph } = require('docx');
const pptxgen = require('pptxgenjs');
const XLSX = require('xlsx');
(async () => {
  const base = process.argv[1];
  const doc = new Document({ sections: [{ children: [new Paragraph('Node DOCX')]}] });
  fs.writeFileSync(base + '.docx', await Packer.toBuffer(doc));
  const book = XLSX.utils.book_new();
  XLSX.utils.book_append_sheet(book, XLSX.utils.aoa_to_sheet([['ready'], ['xlsx']]), 'Sheet1');
  XLSX.writeFile(book, base + '.xlsx');
  const deck = new pptxgen();
  deck.addSlide().addText('Node PPTX', { x: 1, y: 1, w: 4, h: 1 });
  await deck.writeFile({ fileName: base + '.pptx' });
})().catch(error => { console.error(error.stack || error); process.exit(1); });
"""
    _run(["node", "-e", script, str(root / "node-office")], cwd=root, timeout=90)
    expected = {
        "docx": ("node-office.docx", "word/document.xml"),
        "xlsx": ("node-office.xlsx", "xl/workbook.xml"),
        "pptx": ("node-office.pptx", "ppt/presentation.xml"),
    }
    result: dict[str, bool] = {}
    for kind, (filename, member) in expected.items():
        path = root / filename
        with zipfile.ZipFile(path) as archive:
            result[kind] = member in archive.namelist()
    return result


def main() -> int:
    """执行所有 smoke 检查并只输出一条机器可读 JSON。"""
    results: dict[str, object] = {}
    try:
        with tempfile.TemporaryDirectory(prefix="yuxi-offline-smoke-") as directory:
            root = Path(directory)

            document = Document()
            document.add_paragraph("审理报告 {{ subject }}")
            template_path = root / "template.docx"
            document.save(template_path)
            template = DocxTemplate(template_path)
            template.render({"subject": "离线验证"})
            docx_path = root / "report.docx"
            template.save(docx_path)
            assert Document(docx_path).paragraphs[0].text == "审理报告 离线验证"

            workbook = Workbook()
            workbook.active["A1"] = "中文表格"
            xlsx_path = root / "report.xlsx"
            workbook.save(xlsx_path)
            assert load_workbook(xlsx_path).active["A1"].value == "中文表格"

            import xlsxwriter

            writer_path = root / "writer.xlsx"
            with xlsxwriter.Workbook(writer_path) as writer:
                writer.add_worksheet().write("A1", "XlsxWriter ready")
            assert load_workbook(writer_path).active["A1"].value == "XlsxWriter ready"

            import pandas as pd

            assert pd.read_excel(xlsx_path, header=None).iloc[0, 0] == "中文表格"

            presentation = Presentation()
            presentation.slides.add_slide(presentation.slide_layouts[6])
            pptx_path = root / "report.pptx"
            presentation.save(pptx_path)
            assert len(Presentation(pptx_path).slides) == 1

            native_pdf = root / "native.pdf"
            pdf = canvas.Canvas(str(native_pdf))
            pdf.drawString(72, 720, "Offline PDF generation")
            pdf.save()
            assert "Offline PDF generation" in PdfReader(native_pdf).pages[0].extract_text()

            import pypdfium2 as pdfium

            with pdfium.PdfDocument(str(native_pdf)) as rendered_pdf:
                assert len(rendered_pdf) == 1
                assert rendered_pdf[0].render(scale=1).to_pil().width > 0

            import pdfplumber

            with pdfplumber.open(native_pdf) as parsed_pdf:
                assert "Offline PDF generation" in (parsed_pdf.pages[0].extract_text() or "")

            os.environ["MPLCONFIGDIR"] = str(root / "matplotlib-config")
            os.environ["MPLBACKEND"] = "Agg"
            import matplotlib

            matplotlib.use("Agg")
            import matplotlib.pyplot as pyplot

            chart_path = root / "chart.png"
            pyplot.plot([0, 1], [0, 1])
            pyplot.savefig(chart_path)
            pyplot.close()
            assert chart_path.stat().st_size > 0

            _run(
                [
                    "libreoffice",
                    f"-env:UserInstallation={(root / 'lo-profile').as_uri()}",
                    "--headless",
                    "--convert-to",
                    "pdf",
                    "--outdir",
                    str(root),
                    str(docx_path),
                ],
                cwd=root,
            )
            converted = PdfReader(root / "report.pdf")
            converted_text = "".join(page.extract_text() or "" for page in converted.pages)
            assert len(converted.pages) >= 1 and "离线验证" in converted_text
            poppler_text_path = root / "converted.txt"
            _run(
                ["pdftotext", "-layout", str(root / "report.pdf"), str(poppler_text_path)],
                cwd=root,
            )
            assert "离线验证" in poppler_text_path.read_text(encoding="utf-8")

            image_path = root / "ocr.png"
            image = Image.new("RGB", (900, 180), "white")
            english_font = ImageFont.truetype("/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf", 52)
            ImageDraw.Draw(image).text((24, 20), "Yuxi OCR 123", fill="black", font=english_font)
            image.save(image_path)
            ocr_text = _run(["tesseract", str(image_path), "stdout", "-l", "eng"], cwd=root)
            assert "Yuxi" in ocr_text and "123" in ocr_text

            chinese_font_path = Path("/usr/share/fonts/opentype/noto/NotoSansCJK-Regular.ttc")
            assert chinese_font_path.is_file(), "Noto CJK 字体缺失"
            chinese_image_path = root / "ocr-chinese.png"
            chinese_image = Image.new("RGB", (500, 150), "white")
            chinese_font = ImageFont.truetype(str(chinese_font_path), 88)
            ImageDraw.Draw(chinese_image).text((24, 10), "中文测试", fill="black", font=chinese_font)
            chinese_image.save(chinese_image_path)
            chinese_ocr_text = _run(
                ["tesseract", str(chinese_image_path), "stdout", "-l", "chi_sim", "--psm", "7"],
                cwd=root,
            )
            assert "中文" in chinese_ocr_text and "测试" in chinese_ocr_text

            node_files = _write_node_office(root)
            assert all(node_files.values())
            results["python_office"] = {"docx": True, "xlsx": True, "pptx": True, "pdf": True}
            results["office_to_pdf_chinese"] = True
            results["poppler"] = True
            results["tesseract"] = {"english": True, "chinese": True}
            results["node_office"] = node_files

            dotnet_docx = _write_dotnet_docx(root)
            results["dotnet_openxml_offline"] = dotnet_docx
            assert dotnet_docx, "预缓存的 .NET OpenXML 依赖未能离线还原和生成 DOCX"
            assert zipfile.is_zipfile(root / "dotnet-openxml.docx")

            results["uid"] = os.getuid() if hasattr(os, "getuid") else None
            results["writable_tmp"] = True
        print(json.dumps({"ok": True, "checks": results}, ensure_ascii=False, sort_keys=True))
        return 0
    except Exception as exc:  # pragma: no cover - failure output is the container contract.
        print(
            json.dumps(
                {"ok": False, "checks": results, "error": f"{type(exc).__name__}: {exc}"},
                ensure_ascii=False,
                sort_keys=True,
            )
        )
        return 1


if __name__ == "__main__":
    sys.exit(main())
