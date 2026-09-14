"""离线生成并回读办公文件，检查临时用户目录下的系统依赖。"""

import subprocess
import tempfile
from pathlib import Path

from docx import Document
from docxtpl import DocxTemplate
from openpyxl import Workbook, load_workbook
from pptx import Presentation
from pypdf import PdfReader
from reportlab.pdfgen import canvas

with tempfile.TemporaryDirectory() as directory:
    root = Path(directory)
    document = Document()
    document.add_paragraph("审理报告 {{ subject }}")
    document.save(root / "template.docx")
    template = DocxTemplate(root / "template.docx")
    template.render({"subject": "预装环境验证"})
    template.save(root / "report.docx")
    assert Document(root / "report.docx").paragraphs[0].text == "审理报告 预装环境验证"
    workbook = Workbook()
    workbook.active["A1"] = "辅助审理"
    workbook.save(root / "report.xlsx")
    assert load_workbook(root / "report.xlsx").active["A1"].value == "辅助审理"
    slides = Presentation()
    slides.slides.add_slide(slides.slide_layouts[6])
    slides.save(root / "report.pptx")
    assert len(Presentation(root / "report.pptx").slides) == 1
    pdf = canvas.Canvas(str(root / "native.pdf"))
    pdf.drawString(72, 720, "Document runtime ready")
    pdf.save()
    assert "Document runtime ready" in PdfReader(root / "native.pdf").pages[0].extract_text()
    subprocess.run(
        ["libreoffice", f"-env:UserInstallation={root.as_uri()}/lo-profile", "--headless",
         "--convert-to", "pdf", "--outdir", str(root), str(root / "report.docx")],
        check=True, timeout=90,
    )
    converted = PdfReader(root / "report.pdf")
    assert len(converted.pages) >= 1
    assert "预装环境验证" in "".join(page.extract_text() for page in converted.pages)
    print("DOCX/XLSX/PPTX/PDF generation and Chinese Office conversion passed")
