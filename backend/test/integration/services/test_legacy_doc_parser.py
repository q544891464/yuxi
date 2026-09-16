"""用真实 LibreOffice 与统一解析器验证二进制 DOC 正文及表格。"""

import hashlib
import shutil
import subprocess
from pathlib import Path
from types import SimpleNamespace
from unittest.mock import AsyncMock

import pytest
from docx import Document

from yuxi.services.ocr_service import parse_document


pytestmark = pytest.mark.integration


@pytest.mark.asyncio
@pytest.mark.parametrize("extension", ["doc", "wps"])
async def test_binary_doc_preserves_text_table_and_original(tmp_path, monkeypatch, extension):
    """生成真正 Word 97 二进制文件，再通过公开解析入口提取独立已知内容。"""
    executable = shutil.which("soffice") or shutil.which("libreoffice")
    if not executable:
        pytest.skip("需要真实 LibreOffice")
    document = Document()
    document.add_paragraph("合成稽查报告：测试企业收款核对。")
    table = document.add_table(rows=2, cols=2)
    table.cell(0, 0).text = "事项"
    table.cell(0, 1).text = "金额"
    table.cell(1, 0).text = "银行收款"
    table.cell(1, 1).text = "800000.00"
    source = tmp_path / "report.docx"
    document.save(source)
    subprocess.run(
        [
            executable,
            "--headless",
            f"-env:UserInstallation={(tmp_path / 'writer-profile').as_uri()}",
            "--convert-to",
            "doc:MS Word 97",
            "--outdir",
            str(tmp_path),
            str(source),
        ],
        check=True,
        timeout=90,
        capture_output=True,
    )
    # WPS 可保存 Word 兼容内容；此用例不冒充历史 WPS 专有格式夹具。
    legacy = tmp_path / f"案件.{extension.upper()}"
    (tmp_path / "report.doc").rename(legacy)
    original = legacy.read_bytes()
    assert original.startswith(bytes.fromhex("D0CF11E0A1B11AE1")), "夹具必须是旧版二进制 Word"
    markdown = await parse_document(str(legacy), params={"ocr_engine": "disable"})
    assert "合成稽查报告" in markdown
    assert "银行收款" in markdown and "800000.00" in markdown
    assert "事项" in markdown and "金额" in markdown
    assert hashlib.sha256(legacy.read_bytes()).digest() == hashlib.sha256(original).digest()

    # 文件服务边界使用本地磁盘替身，工具到真实解析器及 Markdown 落盘均实际执行。
    from yuxi.agents.toolkits.buildin import tools

    workdir = "/home/gem/user-data/projects/doc-check"
    virtual_source = f"{workdir}/report.{extension}"
    output = tmp_path / "parsed.md"

    def download(path, target, max_bytes):
        """仅提供本用例已经生成的授权文件。"""
        assert path == virtual_source and len(original) < max_bytes
        shutil.copyfile(legacy, target)
        return len(original)

    def upload(path, source):
        """核对工具产物归属并写入真实 Markdown 文件。"""
        assert path == f"{workdir}/outputs/ocr/report.md"
        shutil.copyfile(source, output)

    backend = SimpleNamespace(
        download_authorized_file_to_path=download,
        upload_authorized_file_from_path=upload,
        regular_file_exists=lambda path: False,
    )
    monkeypatch.setattr(tools, "ProvisionerSandboxBackend", lambda **kwargs: backend)
    monkeypatch.setattr(
        tools, "system_options", SimpleNamespace(get=AsyncMock(return_value={"default_ocr_engine": "disable"}))
    )
    context = {
        "thread_id": "doc-test",
        "runtime_scope_id": "doc-test",
        "uid": "doc-user",
        "workdir_relative_path": "projects/doc-check",
        "workdir_path": workdir,
    }
    runtime = SimpleNamespace(config={"configurable": context}, context=SimpleNamespace(**context), state={})
    result = await tools.ocr_parse_file.coroutine(file_path=virtual_source, runtime=runtime)
    saved = Path(output).read_text(encoding="utf-8")
    assert "银行收款" in saved and "800000.00" in saved
    assert result["char_count"] == len(saved)
    assert result["parsed_path"] == f"{workdir}/outputs/ocr/report.md"
    assert legacy.read_bytes() == original
